# -*- coding: utf-8 -*-
"""Generate FacilityIQ_Project_Info.xlsx — links, logins, commands, status."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(r"C:\Users\haris\Downloads\FacilityIQ_Project_Info.xlsx")

HDR = Font(bold=True, color="FFFFFF", size=11)
FILL = PatternFill("solid", fgColor="1F4E5F")
TITLE = Font(bold=True, size=14, color="1F4E5F")


def read_gemini_key() -> str:
    env = ROOT / ".env"
    if env.exists():
        for line in env.read_text(encoding="utf-8").splitlines():
            if line.startswith("GEMINI_API_KEY="):
                return line.split("=", 1)[1].strip()
    return "(see facilityiq/.env)"


def sheet(ws, title, headers, rows, widths):
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws.append([])
    ws.append(headers)
    hr = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(hr, c)
        cell.font = HDR
        cell.fill = FILL
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=hr + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "Links"
    sheet(ws, "FacilityIQ — Links", ["What", "URL / Location"], [
        ("GitHub repo (code+models+data+docs)",
         "https://github.com/subhikshathirunavukkarasu0503/facilityiq  (private)"),
        ("CI runs (GitHub Actions)",
         "https://github.com/subhikshathirunavukkarasu0503/facilityiq/actions"),
        ("Model metrics",
         "https://github.com/subhikshathirunavukkarasu0503/facilityiq/tree/main/models"),
        ("Mid-term report (markdown)", "repo: docs/midterm_report.md"),
        ("Mid-term SUBMISSION doc",
         r"Downloads\S4-I-18_Subhiksha_Thirunavukkarasu_MidTermDoc.docx (also repo docs/)"),
        ("Project handbook", "repo: docs/PROJECT_HANDBOOK.md"),
        ("Portal — cloud", "https://facilityiq-subhiksha.streamlit.app  "
         "(set to PUBLIC in Streamlit sharing settings)"),
        ("Portal — local", "http://localhost:8501"),
        ("Streamlit Cloud console", "https://share.streamlit.io"),
        ("Azure storage (data lake)",
         "portal.azure.com → RG rg-pSiddhi3.0-2026-01-sem4-Subhiksha → "
         "storage fiqlake28091 → container telemetry (196 blobs)"),
        ("Gemini key management", "https://aistudio.google.com/apikey"),
        ("Weather API (keyless)", "https://open-meteo.com"),
        ("Local project folder",
         r"C:\Users\haris\OneDrive\Desktop\Claude Projects\facilityiq"),
    ], [38, 95])

    ws = wb.create_sheet("App Logins")
    sheet(ws, "Portal demo accounts (in-app login)",
          ["Username", "Password", "Role", "Screens visible"], [
              ("viewer", "viewer@123", "Executive Viewer", "Health Overview"),
              ("tech", "tech@123", "Technician",
               "+ Equipment Detail, AI Predictive Intelligence"),
              ("manager", "manager@123", "Facilities Manager",
               "+ Maintenance Scheduling"),
              ("admin", "admin@123", "Administrator",
               "All screens + Admin panel"),
          ], [14, 16, 22, 50])

    ws = wb.create_sheet("Accounts & Keys")
    sheet(ws, "Service accounts & secrets",
          ["Service", "Account", "Auth / where secret lives", "Notes"], [
              ("GitHub", "subhikshathirunavukkarasu0503",
               "gh CLI logged in on this laptop; web login = your password",
               "Repo private; add L&D reviewer as collaborator for evaluation"),
              ("Azure", "subhiksha.thirunavukkarasu@psiog.com",
               "az CLI logged in (device-code); web = your Psiog password",
               "OFFICE subscription — free tier only; teardown script in repo"),
              ("Google AI Studio", "subhikshathirunavukkarasu0503@gmail.com",
               "your Google password", "Manage / revoke API keys here"),
              ("Gemini API key", read_gemini_key(),
               "Local: facilityiq/.env (gitignored) · Cloud: Streamlit app "
               "Settings → Secrets → GEMINI_API_KEY",
               "Revoke/rotate at aistudio.google.com/apikey"),
              ("SECURITY TODO", "—",
               "Rotate Azure + Google passwords (were pasted in chat)",
               "Do this today"),
          ], [22, 46, 58, 55])

    ws = wb.create_sheet("Commands")
    sheet(ws, "Run commands (from facilityiq folder)", ["Task", "Command"], [
        ("Start portal locally",
         r".venv\Scripts\streamlit run src\facilityiq\portal\app.py"),
        ("Regenerate telemetry (14 days)",
         r"set PYTHONPATH=src && .venv\Scripts\python -m facilityiq.simulators.run_simulation"),
        ("Retrain all 3 ML models",
         r"set PYTHONPATH=src && .venv\Scripts\python -m facilityiq.ml.train_all"),
        ("Run full test suite + coverage", r".venv\Scripts\pytest --cov"),
        ("Capture portal screenshots (Selenium)",
         r".venv\Scripts\python scripts\capture_screens.py"),
        ("Delete Azure storage after final demo",
         r"powershell scripts\azure_teardown.ps1"),
    ], [36, 82])

    ws = wb.create_sheet("Status")
    sheet(ws, "Project status (as of 19-Jul-2026)", ["Item", "Status"], [
        ("Mid-term Week-10 checkpoint",
         "~90% — all items working; IoT Hub path pending IT"),
        ("Tests / coverage", "79 passing / 88% measured (CI gate 80%)"),
        ("ML models", "HVAC RF F1 0.997 · Motor GBM F1 0.974 · "
         "Electrical anomaly recall 1.00"),
        ("Cloud", "GitHub full sync · Azure Blob lake 196 blobs · "
         "Streamlit Cloud deployed (make public)"),
        ("Budget", "₹0 spent of ₹2,500"),
        ("Blocked on Psiog IT", "Create F1 IoT Hub + F1 App Service plan in "
         "RG rg-pSiddhi3.0-2026-01-sem4-Subhiksha"),
        ("Phase 2 remaining", "IoT Hub streaming, Functions, Power BI ×3, "
         "Locust load test, Great Expectations, full Selenium E2E"),
    ], [32, 90])

    wb.save(str(OUT))
    print("saved", OUT)


if __name__ == "__main__":
    main()
