# -*- coding: utf-8 -*-
"""Minimal markdown -> docx converter for project docs (headings, tables,
code blocks, bullets, numbered lists, bold)."""

import re
import sys
from pathlib import Path

import docx
from docx.shared import Pt, RGBColor


def add_md_runs(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def convert(src: Path, dst: Path) -> None:
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    lines = src.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x5F)
        elif line.startswith("|") and i + 1 < len(lines) \
                and set(lines[i + 1].replace("|", "").strip()) <= set("-: "):
            head = [c.strip() for c in line.strip("|").split("|")]
            body = []
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                body.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            i -= 1
            table = doc.add_table(rows=1 + len(body), cols=len(head))
            table.style = "Light Grid Accent 1"
            for c, h in enumerate(head):
                cell_p = table.rows[0].cells[c].paragraphs[0]
                add_md_runs(cell_p, h)
                for r in cell_p.runs:
                    r.bold = True
            for r, row in enumerate(body, 1):
                for c, val in enumerate(row[:len(head)]):
                    add_md_runs(table.rows[r].cells[c].paragraphs[0], val)
        elif m := re.match(r"^(#{1,4})\s+(.*)", line):
            doc.add_heading(m.group(2), level=len(m.group(1)))
        elif re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            add_md_runs(p, re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_md_runs(p, re.sub(r"^\s*\d+\.\s+", "", line))
        elif line.strip() == "---":
            pass
        elif line.strip():
            p = doc.add_paragraph()
            add_md_runs(p, line)
        i += 1
    doc.save(str(dst))
    print("saved", dst)


if __name__ == "__main__":
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
